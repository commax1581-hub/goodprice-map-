"""최종 마스터 → 코드화·검색·분석·검수 엑셀 (v2)"""
import json
from urllib.parse import quote
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

SRC = 'data/processed/goodprice_final.csv'
OUT = '착한가격업소_통합데이터_검수_2026-09-20.xlsx'
RESULT_ROWS = 200

m = pd.read_csv(SRC, dtype=str).fillna('')
fac = json.load(open('data/processed/facility_codes.json', encoding='utf-8'))
mapraw = pd.DataFrame(json.load(open('data/raw/goodprice_map_2026-09-19.json', encoding='utf-8'))['items'])

SIDO = [('11', '서울특별시'), ('26', '부산광역시'), ('27', '대구광역시'), ('28', '인천광역시'), ('30', '대전광역시'),
        ('31', '울산광역시'), ('36', '세종특별자치시'), ('41', '경기도'), ('51', '강원특별자치도'), ('43', '충청북도'),
        ('12', '전남광주통합특별시'), ('44', '충청남도'), ('52', '전북특별자치도'), ('47', '경상북도'), ('48', '경상남도'),
        ('50', '제주특별자치도')]
UPJONG = ['한식', '일식', '양식', '중식', '베이커리', '기타요식업', '세탁업', '목욕업', '숙박업', '이용업', '미용업', '기타비요식업']

F = 'Arial'
H_FILL = PatternFill('solid', fgColor='1F4E78'); H_FONT = Font(name=F, bold=True, color='FFFFFF')
IN_FILL = PatternFill('solid', fgColor='FFFF00'); IN_FONT = Font(name=F, color='0000FF', bold=True)
SUB = PatternFill('solid', fgColor='DDEBF7')
BASE = Font(name=F, size=10); BOLD = Font(name=F, bold=True, size=10); TITLE = Font(name=F, bold=True, size=14)
LINK = Font(name=F, size=10, color='0563C1', underline='single')
SMALL = Font(name=F, size=9, color='808080')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font, c.fill, c.border = H_FONT, H_FILL, BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def link_cell(ws, addr, text, url):
    c = ws[addr]; c.value = text; c.hyperlink = url; c.font = LINK


wb = Workbook()

# ============== 업소목록 ==============
ws = wb.active; ws.title = '업소목록'
cols = ['관리번호', '데이터등급', '검수사유', '출처', '시도', '시군구', '행정동', '업종', '세부업종', '업소명', '주소', '정규주소',
        '우편번호', '건물명', '전화번호', '메뉴1', '가격1', '메뉴2', '가격2', '메뉴3', '가격3', '메뉴4', '가격4',
        '최저가격', '위도', '경도', '좌표출처', '주소검증', '사이트업소번호', '카카오장소ID', '카카오장소명', '네이버명',
        '전화검증', '사진'] + [f['code'] for f in fac] + \
       ['공식상세', '카카오', '네이버', '구글', '검색텍스트', '일치', '누적순번']
C = {n: L(i + 1) for i, n in enumerate(cols)}
header(ws, 1, cols)
for f in fac:
    ws[f"{C[f['code']]}1"].comment = Comment(f"{f['label']} (goodprice 엑셀 '{f['column']}' O→Y)", 'build')

N = len(m); LAST = N + 1
NUM = {'가격1', '가격2', '가격3', '가격4'}
FLT = {'위도', '경도'}
for i, r in enumerate(m.itertuples(index=False), start=2):
    d = r._asdict()
    for n in cols:
        if n not in d:
            continue
        c = ws[f'{C[n]}{i}']
        if n in NUM:
            c.value = int(float(d[n])) if d[n] else None
        elif n in FLT:
            c.value = float(d[n]) if d[n] else None
        else:
            c.value = d[n] or None
    ws[f"{C['최저가격']}{i}"] = (f'=IF(COUNT({C["가격1"]}{i},{C["가격2"]}{i},{C["가격3"]}{i},{C["가격4"]}{i})=0,"",'
                              f'MIN({C["가격1"]}{i},{C["가격2"]}{i},{C["가격3"]}{i},{C["가격4"]}{i}))')
    ws[f"{C['검색텍스트']}{i}"] = (f'={C["업소명"]}{i}&" "&{C["메뉴1"]}{i}&" "&{C["메뉴2"]}{i}&" "&'
                               f'{C["메뉴3"]}{i}&" "&{C["메뉴4"]}{i}')
    cond = (f'AND(OR(검색!$C$4="",ISNUMBER(SEARCH(검색!$C$4,{C["검색텍스트"]}{i}))),'
            f'OR(검색!$C$5="",{C["시도"]}{i}=검색!$C$5),OR(검색!$C$6="",{C["업종"]}{i}=검색!$C$6),'
            f'OR(검색!$C$7="",AND(ISNUMBER({C["최저가격"]}{i}),{C["최저가격"]}{i}<=검색!$C$7)),'
            f'OR(검색!$C$8="",{C["데이터등급"]}{i}=검색!$C$8))')
    ws[f"{C['일치']}{i}"] = f'=IF({cond},1,0)'
    ws[f"{C['누적순번']}{i}"] = f'={C["누적순번"]}{i-1}+{C["일치"]}{i}' if i > 2 else f'={C["일치"]}{i}'
    nm, sgg, addr = d['업소명'], d['시군구'], d['주소']
    if d['사진URL']:
        link_cell(ws, f"{C['사진']}{i}", '사진', d['사진URL'])
    if d['사이트업소번호']:
        link_cell(ws, f"{C['공식상세']}{i}", '공식', f"https://www.goodprice.go.kr/bssh/bsshInfo.do?bsshSn={d['사이트업소번호']}")
    if d['카카오링크']:
        link_cell(ws, f"{C['카카오']}{i}", '카카오', d['카카오링크'])
    elif d['위도']:
        link_cell(ws, f"{C['카카오']}{i}", '카카오(위치)', f"https://map.kakao.com/link/map/{quote(nm.replace(',', ' '))},{d['위도']},{d['경도']}")
    link_cell(ws, f"{C['네이버']}{i}", '네이버', f"https://map.naver.com/p/search/{quote(nm + ' ' + sgg)}")
    link_cell(ws, f"{C['구글']}{i}", '구글', f"https://www.google.com/maps/search/?api=1&query={quote(nm + ' ' + addr)}")

for row in ws.iter_rows(min_row=2, max_row=LAST, max_col=len(cols)):
    for c in row:
        if c.font != LINK:
            c.font = BASE
for n in ['가격1', '가격2', '가격3', '가격4', '최저가격']:
    for i in range(2, LAST + 1):
        ws[f'{C[n]}{i}'].number_format = '#,##0'
W = {'관리번호': 9, '데이터등급': 12, '검수사유': 30, '출처': 15, '시도': 14, '시군구': 10, '행정동': 10, '업종': 10,
     '세부업종': 11, '업소명': 22, '주소': 38, '정규주소': 38, '우편번호': 8, '건물명': 14, '전화번호': 13,
     '최저가격': 9, '위도': 11, '경도': 11, '좌표출처': 11, '주소검증': 12, '사이트업소번호': 9,
     '카카오장소ID': 12, '카카오장소명': 18, '네이버명': 18, '전화검증': 14, '검색텍스트': 28}
for n in cols:
    ws.column_dimensions[C[n]].width = W.get(n, 8 if n.startswith('F') or n in ('일치', '누적순번', '사진') else 10)
ws.freeze_panes = 'K2'
ws.auto_filter.ref = f'A1:{L(len(cols))}{LAST}'
for n in ['검색텍스트', '일치', '누적순번']:
    ws.column_dimensions[C[n]].outlineLevel = 1
ws.row_dimensions[1].height = 30


def rng(n):
    return f"업소목록!${C[n]}$2:${C[n]}${LAST}"


# ============== 검수 ==============
chk = m[(m.데이터등급 == '검수필요') | (m.주소검증 == '주소없음')].copy()
cs = wb.create_sheet('검수')
cs['A1'] = f'검수 대상 {len(chk):,}건'; cs['A1'].font = TITLE
cs['A2'] = ('노란 칸(결정·수정주소·참고링크·메모)만 입력하세요. 링크를 눌러 각 플랫폼에서 확인한 뒤 판단합니다. '
            '결정: 승인(정보 맞음) / 수정(주소·상호 변경) / 제외(폐업·중복)'); cs['A2'].font = BASE
ccols = ['관리번호', '결정', '수정주소', '참고링크', '메모', '검수사유', '시도', '시군구', '업종', '업소명', '주소', '전화번호',
         '카카오장소명', '네이버명', '전화검증', '주소검증', '공식', '카카오', '네이버', '구글']
header(cs, 4, ccols)
CC = {n: L(i + 1) for i, n in enumerate(ccols)}
for i, r in enumerate(chk.itertuples(index=False), start=5):
    d = r._asdict()
    for n in ccols:
        if n in d:
            cs[f'{CC[n]}{i}'] = d[n] or None
            cs[f'{CC[n]}{i}'].font = BASE
    for n in ('결정', '수정주소', '참고링크', '메모'):
        c = cs[f'{CC[n]}{i}']; c.fill = IN_FILL; c.font = IN_FONT; c.border = BOX
    nm, sgg, addr = d['업소명'], d['시군구'], d['주소']
    if d['사이트업소번호']:
        link_cell(cs, f"{CC['공식']}{i}", '공식', f"https://www.goodprice.go.kr/bssh/bsshInfo.do?bsshSn={d['사이트업소번호']}")
    if d['위도']:
        link_cell(cs, f"{CC['카카오']}{i}", '카카오', d['카카오링크'] or f"https://map.kakao.com/link/map/{quote(nm.replace(',', ' '))},{d['위도']},{d['경도']}")
    link_cell(cs, f"{CC['네이버']}{i}", '네이버', f"https://map.naver.com/p/search/{quote(nm + ' ' + sgg)}")
    link_cell(cs, f"{CC['구글']}{i}", '구글', f"https://www.google.com/maps/search/?api=1&query={quote(nm + ' ' + addr)}")
dv = DataValidation(type='list', formula1='"승인,수정,제외,보류"', allow_blank=True)
cs.add_data_validation(dv); dv.add(f'B5:B{4 + len(chk)}')
CW = {'관리번호': 9, '결정': 8, '수정주소': 30, '참고링크': 30, '메모': 24, '검수사유': 34, '시도': 13, '시군구': 10,
      '업종': 9, '업소명': 20, '주소': 34, '전화번호': 13, '카카오장소명': 18, '네이버명': 18, '전화검증': 14, '주소검증': 12}
for n in ccols:
    cs.column_dimensions[CC[n]].width = CW.get(n, 8)
cs.freeze_panes = 'F5'
cs.auto_filter.ref = f'A4:{L(len(ccols))}{4 + len(chk)}'

# ============== 코드표 ==============
cd = wb.create_sheet('코드표')
cd['A1'] = '코드표'; cd['A1'].font = TITLE
cd['A2'] = '공식 코드 = goodprice.go.kr 공통코드. ⚠ 표시는 자체 정의 코드(공식 분류 아님).'; cd['A2'].font = BASE
header(cd, 4, ['시도코드(공식)', '시도명', '업소 수'])
for k, (c_, n_) in enumerate(SIDO, start=5):
    cd[f'A{k}'], cd[f'B{k}'] = c_, n_
    cd[f'C{k}'] = f'=COUNTIF({rng("시도")},B{k})'
SIDO_R = (5, 4 + len(SIDO))
header(cd, 4, ['업종(공식)', '업소 수'], col=5)
for k, u in enumerate(UPJONG, start=5):
    cd[f'E{k}'] = u; cd[f'F{k}'] = f'=COUNTIF({rng("업종")},E{k})'
UP_R = (5, 4 + len(UPJONG))
sub = mapraw[['indutyCd', 'indutyNm']].drop_duplicates().sort_values(['indutyCd', 'indutyNm'])
header(cd, 4, ['세부업종코드(공식)', '세부업종명', '업소 수'], col=8)
for k, (c_, n_) in enumerate(sub.itertuples(index=False), start=5):
    cd[f'H{k}'], cd[f'I{k}'] = c_, n_
    cd[f'J{k}'] = f'=COUNTIF({rng("세부업종")},I{k})'
header(cd, 4, ['편의시설 ⚠자체', '항목', '원천 컬럼', 'Y 업소 수'], col=12)
for k, f in enumerate(fac, start=5):
    cd[f'L{k}'], cd[f'M{k}'], cd[f'N{k}'] = f['code'], f['label'], f['column']
    cd[f'O{k}'] = f'=COUNTIF({rng(f["code"])},"Y")'
r0 = 5 + len(fac) + 2
header(cd, r0, ['데이터등급 ⚠자체', '의미'], col=12)
for k, (a_, b_) in enumerate([('확정', '카카오 장소 매칭 또는 전화번호 일치로 확인'),
                              ('확정(네이버)', '카카오에 없으나 네이버 지역검색 100m 이내 동일 상호 확인'),
                              ('검수필요', '세 방법 모두 확인 실패 → 검수 시트 참조')], start=r0 + 1):
    cd[f'L{k}'], cd[f'M{k}'] = a_, b_
r1 = r0 + 5
header(cd, r1, ['주소검증(도로명주소 API)', '의미'], col=12)
for k, (a_, b_) in enumerate([('정상(원문)', '원문 주소 그대로 주소DB에서 확인'),
                              ('정상(정제후)', '상세주소 제거 후 확인'),
                              ('주소없음', '주소DB에 없음 → 원천 데이터 주소 오류 가능')], start=r1 + 1):
    cd[f'L{k}'], cd[f'M{k}'] = a_, b_
for row in cd.iter_rows(min_row=5):
    for c in row:
        if c.value is not None and c.fill != H_FILL:
            c.font = BASE
for col, w in zip('ABCDEFGHIJKLMNO', [12, 18, 9, 3, 14, 9, 3, 14, 14, 9, 3, 20, 52, 20, 10]):
    cd.column_dimensions[col].width = w

# ============== 검색 ==============
ss = wb.create_sheet('검색')
ss['A1'] = '착한가격업소 검색'; ss['A1'].font = TITLE
ss['A2'] = '노란 칸(C4~C8)만 입력하세요. 빈 칸은 조건 없음. 검색어는 업소명+메뉴에서 찾습니다.'; ss['A2'].font = BASE
for k, (a_, b_) in enumerate([('검색어', '예: 김밥'), ('시도', '목록 선택'), ('업종', '목록 선택'),
                              ('최저가격 이하(원)', '예: 5000'), ('데이터등급', '확정 / 확정(네이버) / 검수필요')], start=4):
    ss[f'B{k}'] = a_; ss[f'B{k}'].font = BOLD
    c = ss[f'C{k}']; c.fill = IN_FILL; c.font = IN_FONT; c.border = BOX
    ss[f'D{k}'] = b_; ss[f'D{k}'].font = SMALL
ss['C4'] = '김밥'
d1 = DataValidation(type='list', formula1=f'=코드표!$B${SIDO_R[0]}:$B${SIDO_R[1]}', allow_blank=True)
d2 = DataValidation(type='list', formula1=f'=코드표!$E${UP_R[0]}:$E${UP_R[1]}', allow_blank=True)
d3 = DataValidation(type='list', formula1='"확정,확정(네이버),검수필요"', allow_blank=True)
for dv_, cell in ((d1, 'C5'), (d2, 'C6'), (d3, 'C8')):
    ss.add_data_validation(dv_); dv_.add(cell)
ss['B10'] = '검색 결과 수'; ss['B10'].font = BOLD
ss['C10'] = f'=MAX({rng("누적순번")})'; ss['C10'].font = BOLD; ss['C10'].number_format = '#,##0'
ss['D10'] = f'=IF(C10>{RESULT_ROWS},"※ 상위 {RESULT_ROWS}건만 표시","")'; ss['D10'].font = Font(name=F, size=9, color='C00000')
out = ['관리번호', '데이터등급', '시도', '시군구', '업종', '업소명', '메뉴1', '가격1', '최저가격', '주소', '전화번호', '위도', '경도']
header(ss, 12, ['#', '행'] + out)
for k in range(1, RESULT_ROWS + 1):
    r = 12 + k
    ss[f'A{r}'] = k; ss[f'A{r}'].font = BASE
    ss[f'B{r}'] = f'=IF(A{r}>$C$10,"",MATCH(A{r},{rng("누적순번")},0))'; ss[f'B{r}'].font = SMALL
    for j, n in enumerate(out):
        cl = L(3 + j)
        num = n in ('가격1', '최저가격', '위도', '경도')
        ss[f'{cl}{r}'] = (f'=IF($B{r}="","",IF(INDEX({rng(n)},$B{r})="","",INDEX({rng(n)},$B{r})))' if num
                          else f'=IF($B{r}="","",INDEX({rng(n)},$B{r})&"")')
        ss[f'{cl}{r}'].font = BASE
        if n in ('가격1', '최저가격'):
            ss[f'{cl}{r}'].number_format = '#,##0'
for col, w in zip('ABCDEFGHIJKLMNOP', [5, 6, 10, 12, 13, 10, 9, 20, 18, 9, 9, 34, 13, 11, 11]):
    ss.column_dimensions[col].width = w
ss.freeze_panes = 'A13'

# ============== 분석 ==============
an = wb.create_sheet('분석')
an['A1'] = '착한가격업소 분석 (업소목록 자동 집계)'; an['A1'].font = TITLE
an['A3'] = '1. 시도 × 업종'; an['A3'].font = BOLD
header(an, 4, ['시도'] + UPJONG + ['합계'])
for k, (_, n_) in enumerate(SIDO, start=5):
    an[f'A{k}'] = n_; an[f'A{k}'].font = BOLD
    for j, u in enumerate(UPJONG):
        cl = L(2 + j)
        an[f'{cl}{k}'] = f'=COUNTIFS({rng("시도")},$A{k},{rng("업종")},{cl}$4)'
    an[f'{L(2+len(UPJONG))}{k}'] = f'=SUM(B{k}:{L(1+len(UPJONG))}{k})'
tot = 5 + len(SIDO)
an[f'A{tot}'] = '합계'; an[f'A{tot}'].font = BOLD
for j in range(len(UPJONG) + 1):
    cl = L(2 + j)
    an[f'{cl}{tot}'] = f'=SUM({cl}5:{cl}{tot-1})'; an[f'{cl}{tot}'].font = BOLD; an[f'{cl}{tot}'].fill = SUB

s2 = tot + 2
an[f'A{s2}'] = '2. 업종별 가격(최저가격, 원)'; an[f'A{s2}'].font = BOLD
header(an, s2 + 1, ['업종', '업소 수', '가격 있는 업소', '평균', '최소', '최대'])
for k, u in enumerate(UPJONG, start=s2 + 2):
    an[f'A{k}'] = u
    an[f'B{k}'] = f'=COUNTIF({rng("업종")},A{k})'
    an[f'C{k}'] = f'=COUNTIFS({rng("업종")},A{k},{rng("최저가격")},">0")'
    an[f'D{k}'] = f'=IFERROR(AVERAGEIFS({rng("최저가격")},{rng("업종")},A{k},{rng("최저가격")},">0"),"")'
    an[f'E{k}'] = f'=IF(C{k}=0,"",_xlfn.MINIFS({rng("최저가격")},{rng("업종")},A{k},{rng("최저가격")},">0"))'
    an[f'F{k}'] = f'=IF(C{k}=0,"",_xlfn.MAXIFS({rng("최저가격")},{rng("업종")},A{k},{rng("최저가격")},">0"))'
    for cl in 'DEF':
        an[f'{cl}{k}'].number_format = '#,##0'

s3 = s2 + 2 + len(UPJONG) + 2
an[f'A{s3}'] = '3. 시도별 편의시설 보유율 (분모: 사이트 결합 업소)'; an[f'A{s3}'].font = BOLD
header(an, s3 + 1, ['시도', '분모'] + [f['label'] for f in fac])
for k, (_, n_) in enumerate(SIDO, start=s3 + 2):
    an[f'A{k}'] = n_
    an[f'B{k}'] = f'=COUNTIFS({rng("시도")},A{k},{rng("출처")},"<>공공데이터만")'
    for j, f in enumerate(fac):
        cl = L(3 + j)
        an[f'{cl}{k}'] = f'=IF($B{k}=0,"",COUNTIFS({rng("시도")},$A{k},{rng(f["code"])},"Y")/$B{k})'
        an[f'{cl}{k}'].number_format = '0.0%'

s4 = s3 + 2 + len(SIDO) + 2
an[f'A{s4}'] = '4. 데이터 품질 현황'; an[f'A{s4}'].font = BOLD
header(an, s4 + 1, ['구분', '값', '건수', '비율'])
items = [('데이터등급', '확정'), ('데이터등급', '확정(네이버)'), ('데이터등급', '검수필요'),
         ('좌표출처', 'goodprice'), ('좌표출처', 'kakao_addr'), ('좌표출처', 'kakao_place'),
         ('주소검증', '정상(원문)'), ('주소검증', '정상(정제후)'), ('주소검증', '주소없음'),
         ('출처', '공공데이터+사이트'), ('출처', '공공데이터만'), ('출처', '사이트만')]
for k, (a_, b_) in enumerate(items, start=s4 + 2):
    an[f'A{k}'], an[f'B{k}'] = a_, b_
    an[f'C{k}'] = f'=COUNTIF({rng(a_)},B{k})'
    an[f'D{k}'] = f'=C{k}/COUNTA({rng("관리번호")})'; an[f'D{k}'].number_format = '0.0%'
kr = s4 + 2 + len(items) + 1
an[f'A{kr}'] = '카카오 장소ID 보유'; an[f'A{kr}'].font = BOLD
an[f'C{kr}'] = f'=COUNTIF({rng("카카오장소ID")},"?*")'
an[f'D{kr}'] = f'=C{kr}/COUNTA({rng("관리번호")})'; an[f'D{kr}'].number_format = '0.0%'
an[f'A{kr+1}'] = '공식 사진 보유'; an[f'A{kr+1}'].font = BOLD
an[f'C{kr+1}'] = f'=COUNTIF({rng("사진")},"?*")'
an[f'D{kr+1}'] = f'=C{kr+1}/COUNTA({rng("관리번호")})'; an[f'D{kr+1}'].number_format = '0.0%'
for row in an.iter_rows(min_row=5):
    for c in row:
        if c.value is not None and not c.font.bold and c.fill != H_FILL:
            c.font = BASE
an.column_dimensions['A'].width = 20
for j in range(2, 18):
    an.column_dimensions[L(j)].width = 11

# ============== 안내 ==============
rd = wb.create_sheet('안내', 0)
lines = [('착한가격업소 통합 데이터 (개인 개발 사례용)', TITLE), ('', BASE),
         ('■ 원천 데이터', BOLD),
         ('1) 공공데이터포털 행정안전부_착한가격업소 현황 CSV (기준 2026-06-30, 12,645건)', BASE),
         ('2) goodprice.go.kr 지도 데이터 (2026-09-19 수집, 12,882건) — 좌표·업소번호·세부업종', BASE),
         ('3) goodprice.go.kr 엑셀 다운로드 (2026-09-19, 12,900건) — 편의시설·사진', BASE),
         ('   출처: 행정안전부 착한가격업소(goodprice.go.kr), 공공데이터포털(data.go.kr)', BASE), ('', BASE),
         ('■ 검증에 사용한 API (모두 무료)', BOLD),
         ('· 카카오 로컬 — 주소→좌표 보정, 장소ID 매칭, 전화번호 대조', BASE),
         ('· 네이버 검색(지역) — 카카오 미등록 업소 실재 확인', BASE),
         ('· 행정안전부 도로명주소 검색 API — 주소 검증, 우편번호·행정동 확보', BASE), ('', BASE),
         ('■ 검증 결과', BOLD),
         ('· 좌표 확보 13,118건(100%) — 사이트 12,882 + 카카오 보정 236', BASE),
         ('· 카카오 장소ID 확보 10,494건 / 공식 사진 12,431건', BASE),
         ('· 데이터등급: 확정 10,792 · 확정(네이버) 1,667 · 검수필요 659', BASE),
         ('· 주소검증: 정상 12,877 · 주소없음 241(원천 주소 오류 가능)', BASE), ('', BASE),
         ('■ 시트 구성', BOLD),
         ('업소목록: 1행 1업소. 코드·좌표·검증결과·플랫폼 링크. 오른쪽 접힌 3열은 검색용 계산열', BASE),
         ('검수: 검수필요 + 주소오류 업소. 노란 칸에 결정/수정주소/참고링크/메모 입력', BASE),
         ('코드표 / 검색 / 분석: 코드 정의, 조건 검색, 자동 집계', BASE), ('', BASE),
         ('■ 주의사항', BOLD),
         ('⚠ 시도 명칭은 사이트 현행 기준(광주광역시·전라남도 → 전남광주통합특별시). 원천값은 goodprice_final.csv의 원천시도 열', BASE),
         ('⚠ 편의시설은 사이트 결합 업소만 값 있음(공공데이터만 업소는 N)', BASE),
         ('⚠ 사진·후기는 저장하지 않고 원본 링크로만 연결', BASE),
         ('⚠ 네이버·구글 링크는 검색 링크(장소ID 미확보). 카카오는 장소ID 확보분만 장소 페이지 직결', BASE),
         ('⚠ 폐업 여부는 미확인. 지방행정 인허가 데이터 연계 시 판정 가능', BASE)]
for k, (t, f_) in enumerate(lines, start=1):
    rd[f'A{k}'] = t; rd[f'A{k}'].font = f_
rd.column_dimensions['A'].width = 115

wb.save(OUT)
print('saved', OUT, '| rows', N, '| 검수', len(chk))
