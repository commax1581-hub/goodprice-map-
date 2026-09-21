"""착한가격업소 마스터 → 코드화·검색·분석 엑셀 생성 (openpyxl)"""
import json
from urllib.parse import quote
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

SRC = 'data/processed/goodprice_master.csv'
OUT = '착한가격업소_코드화_검색_분석_2026-09-19.xlsx'
RESULT_ROWS = 200

m = pd.read_csv(SRC, dtype=str).fillna('')
fac = json.load(open('data/processed/facility_codes.json', encoding='utf-8'))
mapraw = pd.DataFrame(json.load(open('data/raw/goodprice_map_2026-09-19.json', encoding='utf-8'))['items'])
report = json.load(open('data/processed/match_report.json', encoding='utf-8'))

SIDO = [('11', '서울특별시'), ('26', '부산광역시'), ('27', '대구광역시'), ('28', '인천광역시'), ('30', '대전광역시'),
        ('31', '울산광역시'), ('36', '세종특별자치시'), ('41', '경기도'), ('51', '강원특별자치도'), ('43', '충청북도'),
        ('12', '전남광주통합특별시'), ('44', '충청남도'), ('52', '전북특별자치도'), ('47', '경상북도'), ('48', '경상남도'),
        ('50', '제주특별자치도')]  # goodprice 공통코드 COM02 순서
UPJONG = ['한식', '일식', '양식', '중식', '베이커리', '기타요식업', '세탁업', '목욕업', '숙박업', '이용업', '미용업', '기타비요식업']

F = 'Arial'
H_FILL = PatternFill('solid', fgColor='1F4E78'); H_FONT = Font(name=F, bold=True, color='FFFFFF')
IN_FILL = PatternFill('solid', fgColor='FFFF00'); IN_FONT = Font(name=F, color='0000FF', bold=True)
SUB_FILL = PatternFill('solid', fgColor='DDEBF7')
BASE = Font(name=F, size=10); BOLD = Font(name=F, bold=True, size=10); TITLE = Font(name=F, bold=True, size=14)
LINK = Font(name=F, size=10, color='0563C1', underline='single')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font, c.fill, c.alignment, c.border = H_FONT, H_FILL, Alignment(horizontal='center', vertical='center', wrap_text=True), BOX


wb = Workbook()

# ================= 업소목록 =================
ws = wb.active; ws.title = '업소목록'
cols = ['관리번호', '출처', '매칭방법', '검수필요', '시도', '원천시도', '시군구', '업종', '세부업종', '업종코드', '업소명', '주소', '전화번호',
        '메뉴1', '가격1', '메뉴2', '가격2', '메뉴3', '가격3', '메뉴4', '가격4', '사이트대표메뉴', '사이트대표가격', '최저가격',
        '위도', '경도', '좌표출처', '사이트업소번호', '이미지수'] + [f['code'] for f in fac] + \
       ['공식상세', '카카오맵', '네이버지도', '구글맵', '검색텍스트', '일치', '누적순번']
C = {name: L(i + 1) for i, name in enumerate(cols)}
header(ws, 1, cols)
for f in fac:
    ws[f"{C[f['code']]}1"].comment = Comment(f"{f['label']} (원천: goodprice 엑셀 '{f['column']}' O→Y)", 'build')

N = len(m); LAST = N + 1
NUM = {'가격1', '가격2', '가격3', '가격4', '사이트대표가격', '이미지수'}
FLT = {'위도', '경도'}
for i, r in enumerate(m.itertuples(index=False), start=2):
    d = r._asdict()
    for name in cols:
        c = ws[f'{C[name]}{i}']
        if name in d and name in NUM:
            c.value = int(float(d[name])) if d[name] not in ('', None) else None
        elif name in d and name in FLT:
            c.value = float(d[name]) if d[name] else None
        elif name in d:
            c.value = d[name] or None
    ws[f"{C['검수필요']}{i}"] = f'=IF(OR(LEFT({C["매칭방법"]}{i},2)="T3",{C["좌표출처"]}{i}="미확보"),"Y","N")'
    ws[f"{C['최저가격']}{i}"] = (f'=IF(COUNT({C["가격1"]}{i},{C["가격2"]}{i},{C["가격3"]}{i},{C["가격4"]}{i},{C["사이트대표가격"]}{i})=0,"",'
                              f'MIN({C["가격1"]}{i},{C["가격2"]}{i},{C["가격3"]}{i},{C["가격4"]}{i},{C["사이트대표가격"]}{i}))')
    ws[f"{C['검색텍스트']}{i}"] = (f'={C["업소명"]}{i}&" "&{C["메뉴1"]}{i}&" "&{C["메뉴2"]}{i}&" "&{C["메뉴3"]}{i}&" "&'
                               f'{C["메뉴4"]}{i}&" "&{C["사이트대표메뉴"]}{i}')
    cond = (f'AND(OR(검색!$C$4="",ISNUMBER(SEARCH(검색!$C$4,{C["검색텍스트"]}{i}))),'
            f'OR(검색!$C$5="",{C["시도"]}{i}=검색!$C$5),OR(검색!$C$6="",{C["업종"]}{i}=검색!$C$6),'
            f'OR(검색!$C$7="",AND(ISNUMBER({C["최저가격"]}{i}),{C["최저가격"]}{i}<=검색!$C$7)),'
            f'OR(검색!$C$8="",{C["좌표출처"]}{i}<>"미확보"))')
    ws[f"{C['일치']}{i}"] = f'=IF({cond},1,0)'
    prev = f'{C["누적순번"]}{i-1}' if i > 2 else '0'
    ws[f"{C['누적순번']}{i}"] = f'={prev}+{C["일치"]}{i}'
    # 링크 (플랫폼 원본 페이지로 이동, 콘텐츠 저장 없음)
    name_q = d['업소명'].replace(',', ' ')
    short = f"{d['업소명']} {d['시군구']}"
    if d['사이트업소번호']:
        c = ws[f"{C['공식상세']}{i}"]; c.value = '공식'; c.hyperlink = f"https://www.goodprice.go.kr/bssh/bsshInfo.do?bsshSn={d['사이트업소번호']}"; c.font = LINK
    if d['위도']:
        c = ws[f"{C['카카오맵']}{i}"]; c.value = '카카오'; c.hyperlink = f"https://map.kakao.com/link/map/{quote(name_q)},{d['위도']},{d['경도']}"; c.font = LINK
    c = ws[f"{C['네이버지도']}{i}"]; c.value = '네이버'; c.hyperlink = f"https://map.naver.com/p/search/{quote(short)}"; c.font = LINK
    c = ws[f"{C['구글맵']}{i}"]; c.value = '구글'; c.hyperlink = f"https://www.google.com/maps/search/?api=1&query={quote(d['업소명'] + ' ' + d['주소'])}"; c.font = LINK

for row in ws.iter_rows(min_row=2, max_row=LAST, max_col=len(cols)):
    for c in row:
        if c.font != LINK:
            c.font = BASE
for name in ['가격1', '가격2', '가격3', '가격4', '사이트대표가격', '최저가격']:
    for i in range(2, LAST + 1):
        ws[f'{C[name]}{i}'].number_format = '#,##0'
widths = {'관리번호': 9, '출처': 15, '매칭방법': 20, '검수필요': 7, '시도': 14, '원천시도': 12, '시군구': 10, '업종': 10, '세부업종': 11,
          '업종코드': 7, '업소명': 22, '주소': 40, '전화번호': 14, '최저가격': 9, '위도': 11, '경도': 11, '좌표출처': 10,
          '사이트업소번호': 9, '검색텍스트': 30}
for name in cols:
    ws.column_dimensions[C[name]].width = widths.get(name, 8 if name.startswith('F') or name in ('일치', '누적순번', '이미지수') else 10)
ws.freeze_panes = 'L2'
ws.auto_filter.ref = f'A1:{L(len(cols))}{LAST}'
for name in ['검색텍스트', '일치', '누적순번']:
    ws.column_dimensions[C[name]].outlineLevel = 1
ws.row_dimensions[1].height = 30


def rng(name):
    return f"업소목록!${C[name]}$2:${C[name]}${LAST}"


# ================= 코드표 =================
cs = wb.create_sheet('코드표')
cs['A1'] = '코드표'; cs['A1'].font = TITLE
cs['A2'] = '공식 코드 = goodprice.go.kr 공통코드/응답값 그대로. 자체 코드(F01~F14, 출처·매칭·좌표출처)는 ⚠ 자체 정의(공식 분류 아님).'; cs['A2'].font = BASE
header(cs, 4, ['시도코드(공식 COM02)', '시도명', '업소 수'])
for k, (cd, nm) in enumerate(SIDO, start=5):
    cs[f'A{k}'], cs[f'B{k}'] = cd, nm
    cs[f'C{k}'] = f'=COUNTIF({rng("시도")},B{k})'
SIDO_R = (5, 4 + len(SIDO))
header(cs, 4, ['업종(대분류, 공식 필터)', '업소 수'], col=5)
for k, u in enumerate(UPJONG, start=5):
    cs[f'E{k}'] = u; cs[f'F{k}'] = f'=COUNTIF({rng("업종")},E{k})'
UP_R = (5, 4 + len(UPJONG))
sub = mapraw[['indutyCd', 'indutyNm']].drop_duplicates().sort_values(['indutyCd', 'indutyNm'])
header(cs, 4, ['세부업종코드(공식)', '세부업종명', '업소 수'], col=8)
for k, (cd, nm) in enumerate(sub.itertuples(index=False), start=5):
    cs[f'H{k}'], cs[f'I{k}'] = cd, nm
    cs[f'J{k}'] = f'=COUNTIF({rng("세부업종")},I{k})'
header(cs, 4, ['편의시설코드 ⚠자체', '항목', '원천 컬럼(goodprice 엑셀)', 'Y 업소 수'], col=12)
for k, f in enumerate(fac, start=5):
    cs[f'L{k}'], cs[f'M{k}'], cs[f'N{k}'] = f['code'], f['label'], f['column']
    cs[f'O{k}'] = f'=COUNTIF({rng(f["code"])},"Y")'
r0 = 5 + len(fac) + 2
header(cs, r0, ['출처 코드 ⚠자체', '의미'], col=12)
for k, (a, b) in enumerate([('공공데이터+사이트', '공공데이터포털 CSV 업소가 goodprice 사이트 업소와 결합됨(좌표·편의시설 보유)'),
                            ('공공데이터만', 'CSV에만 있음 → 좌표 미확보(카카오 지오코딩 대상), 사이트에서 지정취소됐을 가능성'),
                            ('사이트만', '사이트에만 있음 → CSV(2026-06-30) 이후 신규 지정 가능성')], start=r0 + 1):
    cs[f'L{k}'], cs[f'M{k}'] = a, b
r1 = r0 + 5
header(cs, r1, ['매칭방법 ⚠자체', '규칙'], col=12)
for k, (a, b) in enumerate([('T1_업소명+주소', '정규화 업소명 + 도로명·건물번호 완전 일치'),
                            ('T2_업소명+시군구(주소상이)', '업소명 일치 + 같은 시도·시군구 내 유일 (지번/도로명 표기 차이, 이전 등)'),
                            ('T3_주소+업소명유사(x.xx)', '도로명주소 일치 + 업소명 유사도 ≥0.5 유일 → 추정 매칭, 검수 권장'),
                            ('미매칭', '결합 실패')], start=r1 + 1):
    cs[f'L{k}'], cs[f'M{k}'] = a, b
for row in cs.iter_rows(min_row=5):
    for c in row:
        if c.value is not None and c.font != H_FONT:
            c.font = BASE
for col, w in zip('ABCDEFGHIJKLMNO', [12, 18, 9, 3, 16, 9, 3, 12, 14, 9, 3, 22, 50, 22, 10]):
    cs.column_dimensions[col].width = w

# ================= 검색 =================
ss = wb.create_sheet('검색')
ss['A1'] = '착한가격업소 검색'; ss['A1'].font = TITLE
ss['A2'] = '노란 칸(C4~C8)만 입력하세요. 빈 칸은 조건 없음. 검색어는 업소명 + 메뉴(메뉴1~4, 사이트대표메뉴)에서 찾습니다.'; ss['A2'].font = BASE
labels = [('검색어', '예: 김밥'), ('시도', '목록에서 선택'), ('업종', '목록에서 선택'), ('최저가격 이하(원)', '예: 5000'),
          ('좌표 있는 업소만', '아무 값(예: Y) 입력 시 적용')]
for k, (a, b) in enumerate(labels, start=4):
    ss[f'B{k}'] = a; ss[f'B{k}'].font = BOLD
    ss[f'C{k}'].fill = IN_FILL; ss[f'C{k}'].font = IN_FONT; ss[f'C{k}'].border = BOX
    ss[f'D{k}'] = b; ss[f'D{k}'].font = Font(name=F, size=9, color='808080')
ss['C4'] = '김밥'  # 예시 입력값
dv1 = DataValidation(type='list', formula1=f'=코드표!$B${SIDO_R[0]}:$B${SIDO_R[1]}', allow_blank=True)
dv2 = DataValidation(type='list', formula1=f'=코드표!$E${UP_R[0]}:$E${UP_R[1]}', allow_blank=True)
ss.add_data_validation(dv1); ss.add_data_validation(dv2); dv1.add('C5'); dv2.add('C6')
ss['B10'] = '검색 결과 수'; ss['B10'].font = BOLD
ss['C10'] = f'=MAX({rng("누적순번")})'; ss['C10'].font = BOLD; ss['C10'].number_format = '#,##0'
ss['D10'] = f'=IF(C10>{RESULT_ROWS},"※ 상위 {RESULT_ROWS}건만 표시, 조건을 좁혀주세요","")'; ss['D10'].font = Font(name=F, size=9, color='C00000')
out_cols = ['관리번호', '시도', '시군구', '업종', '업소명', '메뉴1', '가격1', '최저가격', '주소', '전화번호', '위도', '경도', '검수필요']
header(ss, 12, ['#', '행'] + out_cols)
for k in range(1, RESULT_ROWS + 1):
    r = 12 + k
    ss[f'A{r}'] = k
    ss[f'B{r}'] = f'=IF(A{r}>$C$10,"",MATCH(A{r},{rng("누적순번")},0))'
    for j, name in enumerate(out_cols):
        cl = L(3 + j)
        ss[f'{cl}{r}'] = f'=IF($B{r}="","",INDEX({rng(name)},$B{r})&"")' if name not in ('가격1', '최저가격', '위도', '경도') \
            else f'=IF($B{r}="","",IF(INDEX({rng(name)},$B{r})="","",INDEX({rng(name)},$B{r})))'
        ss[f'{cl}{r}'].font = BASE
        if name in ('가격1', '최저가격'):
            ss[f'{cl}{r}'].number_format = '#,##0'
    ss[f'A{r}'].font = BASE; ss[f'B{r}'].font = Font(name=F, size=8, color='A6A6A6')
for col, w in zip('ABCDEFGHIJKLMNO', [5, 6, 10, 14, 10, 10, 22, 16, 9, 9, 40, 14, 11, 11, 7]):
    ss.column_dimensions[col].width = w
ss.freeze_panes = 'A13'

# ================= 분석 =================
an = wb.create_sheet('분석')
an['A1'] = '착한가격업소 분석 (업소목록 기준 자동 집계)'; an['A1'].font = TITLE
an['A3'] = '1. 시도 × 업종 업소 수'; an['A3'].font = BOLD
header(an, 4, ['시도'] + UPJONG + ['합계'])
for k, (_, nm) in enumerate(SIDO, start=5):
    an[f'A{k}'] = nm; an[f'A{k}'].font = BOLD
    for j, u in enumerate(UPJONG):
        cl = L(2 + j)
        an[f'{cl}{k}'] = f'=COUNTIFS({rng("시도")},$A{k},{rng("업종")},{cl}$4)'
    an[f'{L(2+len(UPJONG))}{k}'] = f'=SUM(B{k}:{L(1+len(UPJONG))}{k})'
tot = 5 + len(SIDO)
an[f'A{tot}'] = '합계'; an[f'A{tot}'].font = BOLD
for j in range(len(UPJONG) + 1):
    cl = L(2 + j); an[f'{cl}{tot}'] = f'=SUM({cl}5:{cl}{tot-1})'; an[f'{cl}{tot}'].font = BOLD; an[f'{cl}{tot}'].fill = SUB_FILL
an[f'A{tot+1}'] = '※ 전체 업소 수와의 차이 = 시도·업종이 목록 밖인 업소(검증용)'; an[f'A{tot+1}'].font = Font(name=F, size=9, color='808080')
an[f'{L(2+len(UPJONG))}{tot+1}'] = f'=COUNTA({rng("관리번호")})-{L(2+len(UPJONG))}{tot}'

s2 = tot + 3
an[f'A{s2}'] = '2. 업종별 가격 (최저가격 기준, 원)'; an[f'A{s2}'].font = BOLD
header(an, s2 + 1, ['업종', '업소 수', '가격 있는 업소', '평균', '최소', '최대'])
for k, u in enumerate(UPJONG, start=s2 + 2):
    an[f'A{k}'] = u
    an[f'B{k}'] = f'=COUNTIF({rng("업종")},A{k})'
    an[f'C{k}'] = f'=COUNTIFS({rng("업종")},A{k},{rng("최저가격")},">0")'
    an[f'D{k}'] = f'=IFERROR(AVERAGEIFS({rng("최저가격")},{rng("업종")},A{k},{rng("최저가격")},">0"),"")'
    an[f'E{k}'] = f'=IF(C{k}=0,"",_xlfn.MINIFS({rng("최저가격")},{rng("업종")},A{k},{rng("최저가격")},">0"))'
    an[f'F{k}'] = f'=IF(C{k}=0,"",_xlfn.MAXIFS({rng("최저가격")},{rng("업종")},A{k},{rng("최저가격")},">0"))'
    for cl in 'DEF':
        an[f'{cl}{k}'].number_format = '#,##0'

s3 = s2 + 2 + len(UPJONG) + 2
an[f'A{s3}'] = '3. 시도별 편의시설 보유율 (분모: 사이트 결합 업소 = 출처가 "공공데이터만"이 아닌 업소)'; an[f'A{s3}'].font = BOLD
header(an, s3 + 1, ['시도', '분모'] + [f['label'] for f in fac])
for k, (_, nm) in enumerate(SIDO, start=s3 + 2):
    an[f'A{k}'] = nm
    an[f'B{k}'] = f'=COUNTIFS({rng("시도")},A{k},{rng("출처")},"<>공공데이터만")'
    for j, f in enumerate(fac):
        cl = L(3 + j)
        an[f'{cl}{k}'] = f'=IF($B{k}=0,"",COUNTIFS({rng("시도")},$A{k},{rng(f["code"])},"Y")/$B{k})'
        an[f'{cl}{k}'].number_format = '0.0%'

s4 = s3 + 2 + len(SIDO) + 2
an[f'A{s4}'] = '4. 데이터 결합 현황'; an[f'A{s4}'].font = BOLD
header(an, s4 + 1, ['구분', '값', '건수'])
items = [('출처', '공공데이터+사이트'), ('출처', '공공데이터만'), ('출처', '사이트만'),
         ('매칭방법', 'T1*'), ('매칭방법', 'T2*'), ('매칭방법', 'T3*'), ('매칭방법', '미매칭'),
         ('좌표출처', 'goodprice'), ('좌표출처', '미확보'), ('검수필요', 'Y')]
for k, (a, b) in enumerate(items, start=s4 + 2):
    an[f'A{k}'], an[f'B{k}'] = a, b
    an[f'C{k}'] = f'=COUNTIF({rng(a)},B{k})'
for row in an.iter_rows(min_row=5):
    for c in row:
        if c.value is not None and c.font not in (H_FONT, BOLD) and not c.font.bold:
            c.font = BASE
an.column_dimensions['A'].width = 18
for j in range(2, 17):
    an.column_dimensions[L(j)].width = 11

# ================= 안내 =================
rd = wb.create_sheet('안내', 0)
lines = [
    ('착한가격업소 코드화·검색·분석 (개인 개발 사례용)', TITLE),
    ('', BASE),
    ('■ 원천 데이터', BOLD),
    ('1) 공공데이터포털 「행정안전부_착한가격업소 현황」 CSV (기준일 2026-06-30, 12,645건) — 기준 목록, 메뉴1~4·가격1~4', BASE),
    ('2) goodprice.go.kr 지도 데이터 (2026-09-19 수집, 시도별 16회 요청, 12,882건) — 위도·경도, 사이트업소번호, 세부업종', BASE),
    ('3) goodprice.go.kr "엑셀 다운로드" (2026-09-19, 12,900건) — 편의시설·지역화폐 Y/N, 이미지 수, 대표메뉴', BASE),
    ('   출처: 행정안전부 착한가격업소(goodprice.go.kr), 공공데이터포털(data.go.kr)', BASE),
    ('', BASE),
    ('■ 시트 구성', BOLD),
    ('업소목록: 전체 업소 1행 1업소, 코드 컬럼·좌표·플랫폼 링크(공식/카카오/네이버/구글) 포함. 오른쪽 접힌 3열은 검색용 계산열', BASE),
    ('코드표: 시도·업종·세부업종(공식 코드) + 편의시설·출처·매칭방법(⚠ 자체 코드)', BASE),
    ('검색: 노란 칸에 조건 입력 → 상위 200건 표시', BASE),
    ('분석: 시도×업종, 업종별 가격, 시도별 편의시설 보유율, 결합 현황 (모두 수식, 업소목록 수정 시 자동 갱신)', BASE),
    ('', BASE),
    ('■ 주의사항', BOLD),
    ('⚠ 시도 "현행화": CSV의 광주광역시·전라남도는 사이트 현행 명칭 "전남광주통합특별시"로 통일 (원천값은 "원천시도" 열에 보존)', BASE),
    ('⚠ 매칭방법 T3은 업소명 유사도 기반 추정 결합 → 검수필요=Y', BASE),
    ('⚠ 좌표출처 "미확보" 업소는 지도 표시 불가 → 카카오 주소검색으로 보정 예정', BASE),
    ('⚠ 편의시설은 사이트 결합 업소만 값이 있음(공공데이터만 업소는 N으로 채워짐, 보유율 분모에서 제외)', BASE),
    ('⚠ 이용후기·업소 사진 원본은 수집하지 않음(링크로만 연결). 플랫폼 링크는 검색/좌표 기반이며 장소ID 매칭 전 단계', BASE),
    ('', BASE),
    ('■ 결합 결과 (수집 시점 스냅샷)', BOLD),
]
for k, (t, fnt) in enumerate(lines, start=1):
    rd[f'A{k}'] = t; rd[f'A{k}'].font = fnt
k0 = len(lines) + 1
snap = [('공공데이터 CSV', report['inputs']['datagokr_csv']), ('사이트 엑셀', report['inputs']['site_excel']),
        ('사이트 지도(좌표)', report['inputs']['site_map']), ('마스터 업소 수(현재 시트)', '=COUNTA(업소목록!A:A)-1')]
for k, (a, b) in enumerate(snap, start=k0):
    rd[f'A{k}'] = a; rd[f'B{k}'] = b; rd[f'A{k}'].font = BASE; rd[f'B{k}'].font = BASE; rd[f'B{k}'].number_format = '#,##0'
rd.column_dimensions['A'].width = 120; rd.column_dimensions['B'].width = 12

wb.save(OUT)
print('saved', OUT, 'rows', N)
